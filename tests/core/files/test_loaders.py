"""ADR-008 E3 (#1070) — more document formats, as loaders at ingestion time.

What each block below pins, and why it needs pinning:

* **one list** — the supported extensions are the registry's. Uploads, the
  knowledge ingest, the CLI and the web picker all used to hold their own copy,
  and the copies had drifted (the ingest globbed `.pdf` separately and read it
  with a weaker reader);
* **every format reads** — fixtures are REAL files built here with the same
  libraries a user's Office would produce them with, not mocks of the loaders;
* **the defences** — magic bytes, the zip-bomb cap, the output cap, XXE, the
  excluded extensions: each is the test that goes red when the defence is
  removed (see the mutation notes in the E3 report);
* **the two consumers** — the knowledge ingest gains the formats and the B026
  PDF pipeline; an upload of a .docx lands chunked in `user_knowledge`.
"""
from __future__ import annotations

import asyncio
import json
import struct
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from fastapi import HTTPException

from core.files import loaders
from core.files.handler import FileHandler
from core.files.loaders import office

SECRET = "TOPSECRET-E3-9f31"


@pytest.fixture()
def fh(tmp_path):
    return FileHandler(tmp_path / "uploads")


# ── fixture builders (real files) ────────────────────────────────

def make_docx(path: Path, sentence: str = "La clau del projecte és Mirmidó-42.") -> Path:
    import docx

    doc = docx.Document()
    doc.add_heading("Informe trimestral", level=1)
    doc.add_paragraph(sentence)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Ciutat"
    table.cell(0, 1).text = "Preu"
    table.cell(1, 0).text = "Girona"
    table.cell(1, 1).text = "12"
    doc.add_paragraph("Paràgraf després de la taula.")
    doc.save(path)
    return path


def make_xlsx(path: Path, rows: int = 2) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendes"
    ws.append(["Producte", "Unitats"])
    for i in range(rows - 1):
        ws.append([f"Pera{i}", 7 + i])
    wb.create_sheet("Buit")
    other = wb.create_sheet("Stock")
    other["A1"] = "Magatzem Nord"
    wb.save(path)
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Pla estratègic"
    s1.placeholders[1].text = "Obrir a Lleida"
    s1.notes_slide.notes_text_frame.text = "Nota del ponent"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "Any"
    tbl.cell(0, 1).text = "Ingressos"
    prs.save(path)
    return path


_CONTAINER = (
    '<?xml version="1.0"?><container version="1.0" '
    'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
    '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
    '</rootfiles></container>'
)


def make_epub(path: Path) -> Path:
    """Hand-built: the spine order (c1, c2) is the REVERSE of the archive order."""
    opf = (
        '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
        '<manifest>'
        '<item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="c1" href="text/ch%201.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="css" href="s.css" media-type="text/css"/>'
        '</manifest><spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>'
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", _CONTAINER)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/ch2.xhtml", "<html><body><p>Capítol dos</p><script>alert(1)</script></body></html>")
        z.writestr(
            "OEBPS/text/ch 1.xhtml",
            '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
            "<head><title>T</title><style>p{color:red}</style></head>"
            "<body><h1>Capítol u</h1><p>Primer&#160;paràgraf.</p><p>Segon.</p></body></html>",
        )
        z.writestr("OEBPS/s.css", "p{}")
    return path


# ── one list ─────────────────────────────────────────────────────

E3_FORMATS = {".docx", ".xlsx", ".pptx", ".epub", ".csv", ".json", ".xml", ".yaml", ".yml", ".toml", ".py", ".ts"}
PRE_E3 = {".txt", ".md", ".markdown", ".text", ".pdf"}


def test_every_consumer_reads_the_registry_list():
    from core.files import handler
    from core.ingest import ingest_knowledge
    from plugins.web_ui_module.api.routes_static import upload_accept

    registry = frozenset(loaders.LOADER_REGISTRY)
    assert E3_FORMATS | PRE_E3 <= registry
    assert handler.SUPPORTED_EXTENSIONS == registry
    assert ingest_knowledge.SUPPORTED_EXTENSIONS == registry
    accepted = set(upload_accept().split(","))
    assert registry <= accepted
    assert accepted - registry == {"image/jpeg", "image/png", "image/webp"}


def test_validate_file_asks_the_registry_not_a_copy(fh, monkeypatch):
    """A loader registered at runtime is accepted by the upload door at once."""
    ok, msg = fh.validate_file("probe.e3probe", 10, b"hello")
    assert ok is False and "Unsupported format" in msg
    monkeypatch.setitem(loaders.LOADER_REGISTRY, ".e3probe", loaders.Loader(func=lambda p: "x"))
    ok, msg = fh.validate_file("probe.e3probe", 10, b"hello")
    assert ok is True, msg


def test_the_web_picker_is_rendered_from_the_registry():
    from plugins.web_ui_module.api.routes_static import register_static_routes
    from fastapi import APIRouter

    ui_dir = Path(__file__).resolve().parents[3] / "plugins" / "web_ui_module" / "ui"
    raw = (ui_dir / "index.html").read_text(encoding="utf-8")
    assert 'id="fileInput" accept="{{NEXE_UPLOAD_ACCEPT}}"' in raw, (
        "the picker must not carry a hand-written copy of the format list"
    )
    router = APIRouter()
    register_static_routes(router, module_ref=SimpleNamespace(ui_dir=ui_dir))
    serve_ui = next(r.endpoint for r in router.routes if getattr(r, "path", "") == "/")
    html = asyncio.run(serve_ui(i18n=None)).body.decode()
    assert "{{NEXE_UPLOAD_ACCEPT}}" not in html
    accept = html.split('id="fileInput" accept="', 1)[1].split('"', 1)[0]
    assert set(accept.split(",")) >= set(loaders.LOADER_REGISTRY)
    assert ".docx" in accept.split(",")


# ── every format reads ───────────────────────────────────────────

def test_docx_paragraphs_and_tables_in_order(tmp_path):
    text = loaders.extract_text(make_docx(tmp_path / "a.docx"))
    assert text == (
        "Informe trimestral\n\nLa clau del projecte és Mirmidó-42.\n\n"
        "Ciutat | Preu\nGirona | 12\n\nParàgraf després de la taula."
    )


def test_xlsx_every_non_empty_sheet(tmp_path):
    text = loaders.extract_text(make_xlsx(tmp_path / "a.xlsx"))
    assert text == "[Sheet: Vendes]\nProducte | Unitats\nPera0 | 7\n\n[Sheet: Stock]\nMagatzem Nord"


def test_pptx_text_tables_and_notes(tmp_path):
    text = loaders.extract_text(make_pptx(tmp_path / "a.pptx"))
    assert text == (
        "[Slide 1]\nPla estratègic\nObrir a Lleida\nNotes: Nota del ponent\n\n"
        "[Slide 2]\nAny | Ingressos"
    )


def test_epub_follows_the_spine_and_drops_scripts(tmp_path):
    text = loaders.extract_text(make_epub(tmp_path / "a.epub"))
    assert text == "Capítol u\nPrimer paràgraf.\nSegon.\n\nCapítol dos"


def test_csv_semicolon_rows_keep_their_column_names(tmp_path):
    p = tmp_path / "a.csv"
    p.write_text("nom;edat\nAnna;30\n\nBernat;41;extra\n", encoding="utf-8")
    assert loaders.extract_text(p) == (
        "Row 1:\n  nom: Anna\n  edat: 30\n\nRow 3:\n  nom: Bernat\n  edat: 41\n  col3: extra"
    )


def test_json_pretty_printed_and_invalid_json_kept_raw(tmp_path):
    p = tmp_path / "a.json"
    p.write_text('{"ciutat":"Vic","n":[1,2]}', encoding="utf-8")
    assert json.loads(loaders.extract_text(p)) == {"ciutat": "Vic", "n": [1, 2]}
    assert loaders.extract_text(p).count("\n") > 3
    bad = tmp_path / "b.json"
    bad.write_text("{not json", encoding="utf-8")
    assert loaders.extract_text(bad) == "{not json"


def test_yaml_and_toml_are_indexed_as_written_comments_included(tmp_path):
    y = tmp_path / "c.yaml"
    y.write_text("# el port del servei\nport: 9119\n", encoding="utf-8")
    assert loaders.extract_text(y) == "# el port del servei\nport: 9119\n"
    t = tmp_path / "c.toml"
    t.write_text("[a]\n# nota\nb = 1\n", encoding="utf-8")
    assert loaders.extract_text(t) == "[a]\n# nota\nb = 1\n"


def test_code_is_text(tmp_path):
    p = tmp_path / "m.py"
    p.write_text("def f():\n    return 'hola'\n", encoding="utf-8")
    assert loaders.extract_text(p) == "def f():\n    return 'hola'\n"


def test_xml_pretty_printed(tmp_path):
    p = tmp_path / "a.xml"
    p.write_text("<r><x>1</x><y>dos</y></r>", encoding="utf-8")
    assert loaders.extract_text(p) == "<r>\n  <x>1</x>\n  <y>dos</y>\n</r>"


# ── magic bytes ──────────────────────────────────────────────────

@pytest.mark.parametrize("ext", [".docx", ".xlsx", ".pptx", ".epub"])
def test_zip_format_that_is_really_text_is_refused(fh, ext):
    ok, msg = fh.validate_file(f"x{ext}", 20, b"just some plain text")
    assert ok is False
    assert msg == f"File content does not match {ext} format"


def test_real_docx_passes_the_magic_check(fh, tmp_path):
    content = make_docx(tmp_path / "a.docx").read_bytes()
    assert fh.validate_file("a.docx", len(content), content) == (True, "")


def test_fake_docx_refused_at_the_door_and_not_saved(fh, tmp_path):
    session_mgr = MagicMock()
    with pytest.raises(HTTPException) as exc:
        asyncio.run(_attach(fh, session_mgr, "informe.docx", b"not a zip at all"))
    assert exc.value.status_code == 400
    assert exc.value.detail == "File content does not match .docx format"
    assert not list(fh._iter_files())
    session_mgr.get_or_create_session.assert_not_called()


def test_zip_with_a_prefix_is_refused_although_zipfile_would_open_it(tmp_path):
    """`zipfile` finds the central directory from the END, so a real docx with
    anything glued in front (a polyglot, a self-extractor) opens fine. The
    magic check at the loader is what refuses it in the knowledge folder,
    where no door looked at the first bytes."""
    good = make_docx(tmp_path / "good.docx").read_bytes()
    p = tmp_path / "polyglot.docx"
    p.write_bytes(b"MZ\x90\x00" + b"\0" * 60 + good)
    assert zipfile.is_zipfile(p)
    with pytest.raises(loaders.LoaderError, match="is not a zip archive"):
        loaders.extract_text(p)


def test_fake_docx_in_the_knowledge_folder_is_skipped_not_crashed(tmp_path):
    """No door checked the magic there: the loader itself does."""
    from core.ingest.ingest_knowledge import read_file

    p = tmp_path / "fals.docx"
    p.write_text("plain text pretending", encoding="utf-8")
    with pytest.raises(loaders.LoaderError, match="not a zip"):
        loaders.extract_text(p)
    assert read_file(p) == ""


# ── UTF-8 rule for text formats ──────────────────────────────────

def test_utf8_char_straddling_byte_4096_is_valid(fh):
    content = b"a" * 4095 + "à".encode() + b" fi"
    assert fh.validate_file("x.txt", len(content), content) == (True, "")


def test_invalid_utf8_after_the_first_4k_is_refused(fh):
    content = b"a" * 5000 + b"\xff"
    ok, msg = fh.validate_file("x.csv", len(content), content)
    assert (ok, msg) == (False, "File content is not valid UTF-8 text")


# ── zip bomb ─────────────────────────────────────────────────────

def _zip_with_member(path: Path, size: int, extra_members: int = 0) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("big.bin", b"\0" * size)
        for i in range(extra_members):
            z.writestr(f"m{i}.txt", b"")
    return path


def test_zip_cap_boundary(tmp_path, monkeypatch):
    monkeypatch.setattr(office, "MAX_ZIP_UNCOMPRESSED", 100_000)
    office.check_zip(_zip_with_member(tmp_path / "at.zip", 100_000))  # at the cap: fine
    with pytest.raises(loaders.LoaderError, match="inflates to 100001"):
        office.check_zip(_zip_with_member(tmp_path / "over.zip", 100_001))


def test_zip_member_count_boundary(tmp_path, monkeypatch):
    monkeypatch.setattr(office, "MAX_ZIP_MEMBERS", 5)
    office.check_zip(_zip_with_member(tmp_path / "at.zip", 1, extra_members=4))
    with pytest.raises(loaders.LoaderError, match="6 archive members"):
        office.check_zip(_zip_with_member(tmp_path / "over.zip", 1, extra_members=5))


def test_default_zip_cap_is_twenty_uploads():
    assert office.MAX_ZIP_UNCOMPRESSED == 20 * 10 * 1024 * 1024
    assert office.MAX_ZIP_MEMBERS == 10_000


def test_bombed_docx_never_reaches_the_parser(tmp_path, monkeypatch, fh):
    """A real .docx with a member that inflates past the cap."""
    src = make_docx(tmp_path / "ok.docx")
    declared = sum(i.file_size for i in zipfile.ZipFile(src).infolist())
    monkeypatch.setattr(office, "MAX_ZIP_UNCOMPRESSED", declared + 1000)
    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/media/pad.bin", b"\0" * 1001)
    import docx

    with patch.object(docx, "Document", side_effect=AssertionError("parser reached")):
        with pytest.raises(loaders.LoaderError, match="inflates"):
            loaders.extract_text(bomb)
    assert fh.extract_text(bomb) == ""


def test_a_member_cannot_inflate_past_its_declared_size(tmp_path):
    """The cap trusts declared sizes; this is why that is sound."""
    p = tmp_path / "liar.zip"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("x.xml", b"A" * 5000)
    data = bytearray(p.read_bytes())
    # Patch the uncompressed size in the local header (offset 22) and the
    # central directory entry (offset 24 from its signature) down to 10.
    struct.pack_into("<I", data, 22, 10)
    cd = data.find(b"PK\x01\x02")
    struct.pack_into("<I", data, cd + 24, 10)
    p.write_bytes(bytes(data))
    office.check_zip(p)  # declares 10 bytes: under any cap
    with zipfile.ZipFile(p) as z, pytest.raises(zipfile.BadZipFile):
        z.read("x.xml")


# ── output cap ───────────────────────────────────────────────────

def test_text_budget_boundary(monkeypatch):
    monkeypatch.setattr(loaders, "MAX_EXTRACTED_CHARS", 10)
    b = loaders.TextBudget(sep="\n")
    b.add("abcd")  # 4 + 1
    b.add("efgh")  # 4 + 1 -> 10: at the cap
    assert b.text() == "abcd\nefgh"
    with pytest.raises(loaders.LoaderError, match="exceeds 10"):
        b.add("")


def test_huge_spreadsheet_stops_at_the_output_cap(tmp_path, monkeypatch):
    p = make_xlsx(tmp_path / "big.xlsx", rows=500)
    full = loaders.extract_text(p)
    monkeypatch.setattr(loaders, "MAX_EXTRACTED_CHARS", len(full) // 2)
    with pytest.raises(loaders.LoaderError, match="exceeds"):
        loaders.extract_text(p)


def test_output_cap_reuses_the_upload_cap():
    from core.files.handler import MAX_FILE_SIZE

    assert loaders.MAX_EXTRACTED_CHARS == MAX_FILE_SIZE == 10 * 1024 * 1024


# ── XML: XXE and billion laughs ──────────────────────────────────

def test_xml_external_entity_is_not_resolved(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text(SECRET, encoding="utf-8")
    p = tmp_path / "evil.xml"
    p.write_text(
        f'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY xxe SYSTEM "{secret.as_uri()}">]>'
        "<r><a>&xxe;</a></r>",
        encoding="utf-8",
    )
    text = loaders.extract_text(p)
    assert SECRET not in text
    assert "&xxe;" in text


def test_the_xml_parser_never_loads_an_external_dtd(tmp_path):
    """`load_dtd=False`, observed: nothing in the loader's OUTPUT changes when
    a DTD is loaded while entities stay unresolved, so the parser is asked."""
    from lxml import etree
    from core.files.loaders.data import safe_xml_parser

    dtd = tmp_path / "ext.dtd"
    dtd.write_text('<!ENTITY e "FROM-DTD">', encoding="utf-8")
    doc = f'<?xml version="1.0"?><!DOCTYPE r SYSTEM "{dtd.as_uri()}"><r>&e;</r>'.encode()
    root = etree.fromstring(doc, safe_xml_parser())
    assert root.getroottree().docinfo.externalDTD is None
    assert b"FROM-DTD" not in etree.tostring(root)


def test_xml_billion_laughs_is_not_expanded(tmp_path):
    ents = '<!ENTITY a0 "lol">' + "".join(
        f'<!ENTITY a{i} "{("&a" + str(i - 1) + ";") * 10}">' for i in range(1, 10)
    )
    p = tmp_path / "lol.xml"
    p.write_text(f'<?xml version="1.0"?><!DOCTYPE r [{ents}]><r>&a9;</r>', encoding="utf-8")
    text = loaders.extract_text(p)
    assert len(text) < 1000
    assert "lollol" not in text


def test_epub_package_xml_external_entity_is_not_resolved(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text(SECRET, encoding="utf-8")
    opf = (
        f'<?xml version="1.0"?><!DOCTYPE package [<!ENTITY xxe SYSTEM "{secret.as_uri()}">]>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0"><manifest>'
        '<item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
        '<spine><itemref idref="c1"/></spine><metadata>&xxe;</metadata></package>'
    )
    p = tmp_path / "evil.epub"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("META-INF/container.xml", _CONTAINER)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/c1.xhtml", "<html><body><p>Text del llibre</p></body></html>")
    assert loaders.extract_text(p) == "Text del llibre"


# ── excluded extensions ──────────────────────────────────────────

@pytest.mark.parametrize("name", [
    "prod.env", "model.pkl", "model.pickle", "settings.ini", "nginx.conf",
    "setup.cfg", "app.config", "db.properties", ".env", "deploy.sh",
])
def test_excluded_and_unlisted_names_are_refused(fh, name):
    ok, msg = fh.validate_file(name, 10, b"KEY=value")
    assert ok is False and msg.startswith("Unsupported format")


def test_the_registry_never_holds_an_excluded_extension():
    assert not (set(loaders.LOADER_REGISTRY) & loaders.EXCLUDED_EXTENSIONS)
    assert {".env", ".pkl", ".pickle", ".ini", ".conf", ".properties"} <= loaders.EXCLUDED_EXTENSIONS


@pytest.mark.parametrize("ext", [".env", ".PKL", ".ini"])
def test_registering_an_excluded_extension_raises(ext):
    with pytest.raises(ValueError, match="excluded"):
        loaders.register_loader([ext])(lambda p: "")
    assert ext.lower() not in loaders.LOADER_REGISTRY


def test_registering_twice_raises():
    with pytest.raises(ValueError, match="already has a loader"):
        loaders.register_loader([".docx"])(lambda p: "")
    assert loaders.LOADER_REGISTRY[".docx"].func is office.load_docx_file


def test_the_sensitive_content_denylist_covers_the_new_text_formats(fh):
    session_mgr = MagicMock()
    with pytest.raises(HTTPException) as exc:
        asyncio.run(_attach(
            fh, session_mgr, "deploy.py",
            b"KEY = '''\n-----BEGIN OPENSSH PRIVATE KEY-----\nabc\n'''\n",
        ))
    assert exc.value.status_code == 400
    assert "sensitive pattern" in exc.value.detail


# ── the knowledge ingest uses the registry ───────────────────────

def test_ingest_discovers_the_new_formats_and_skips_hidden_folders(tmp_path):
    from core.ingest.ingest_knowledge import _discover_documents

    make_docx(tmp_path / "a.docx")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.csv").write_text("x,y\n1,2\n", encoding="utf-8")
    (tmp_path / "c.md").write_text("# c", encoding="utf-8")
    (tmp_path / ".embeddings").mkdir()
    (tmp_path / ".embeddings" / "manifest.json").write_text("{}", encoding="utf-8")
    (tmp_path / "d.exe").write_bytes(b"MZ")
    found = [p.relative_to(tmp_path).as_posix() for p in _discover_documents(tmp_path)]
    assert found == ["a.docx", "c.md", "sub/b.csv"]


def test_ingest_builds_chunks_from_a_docx(tmp_path):
    from core.ingest.ingest_knowledge import _build_file_items

    p = make_docx(tmp_path / "manual.docx", sentence="El codi secret és Tramuntana.")
    items, _coll, n = _build_file_items(p, MagicMock(), "nexe_documentation", lambda m: None, [0])
    assert n >= 1
    assert "El codi secret és Tramuntana." in items[0]["text"]
    assert items[0]["metadata"]["source"] == "manual.docx"


def _fake_pdf_page(default_text, layout_text=None):
    page = MagicMock()
    page.extract_text = MagicMock(
        side_effect=lambda extraction_mode=None: (
            layout_text if extraction_mode == "layout" and layout_text is not None else default_text
        )
    )
    return page


def test_ingest_pdf_gets_the_b026_pipeline(tmp_path):
    """The knowledge reader used to be `text += page.extract_text() + "\\n"`:
    no NFKC (ligatures stayed), no layout retry for glued pages."""
    from core.ingest.ingest_knowledge import read_file

    prose = ("Els agents d'IA local permeten treballar sense enviar dades a tercers. " * 4).strip()
    glued = prose.replace(" ", "")
    reader = MagicMock()
    reader.pages = [_fake_pdf_page("la conﬁguració " + prose), _fake_pdf_page(glued, layout_text=prose)]
    p = tmp_path / "m.pdf"
    p.write_bytes(b"%PDF-1.4")
    with patch("pypdf.PdfReader", return_value=reader):
        text = read_file(p)
    assert "la configuració" in text and "ﬁ" not in text
    assert glued not in text
    assert text.count("Els agents d'IA local") == 8


# ── end to end: an uploaded .docx lands in user_knowledge ────────

async def _attach(fh, session_mgr, filename, content, helper=None):
    from core.files.attach import attach_to_session

    helper = helper or MagicMock(save_document_chunks=AsyncMock(return_value={"success": True, "chunks_saved": 1}))
    with patch("core.memory_facts.helper_for", return_value=helper):
        return await attach_to_session(
            app_state=MagicMock(), session_mgr=session_mgr, file_handler=fh,
            filename=filename, content=content, session_id="e3-session",
        )


def test_uploaded_docx_is_chunked_into_user_knowledge(fh, tmp_path):
    import core.memory_facts.helper as mh_module
    from core.memory_access import KNOWLEDGE_COLLECTION
    from core.memory_facts.helper import MemoryHelper

    mem = MagicMock()
    mem.collection_exists = AsyncMock(return_value=True)
    mem.store_batch = AsyncMock()
    mem.ingest_config = MagicMock(store_batch_size=50)
    helper = MemoryHelper()
    helper._memory_api = mem
    session = MagicMock(id="e3-session")
    session_mgr = MagicMock()
    session_mgr.get_or_create_session.return_value = session

    content = make_docx(tmp_path / "src.docx", sentence="La contrasenya del wifi és Pedraforca-77.").read_bytes()
    original = mh_module._memory_api_instance
    mh_module._memory_api_instance = mem
    try:
        body = asyncio.run(_attach(fh, session_mgr, "informe.docx", content, helper=helper))
    finally:
        mh_module._memory_api_instance = original

    assert body["ingested"] is True and body["chunks"] >= 1
    assert mem.store_batch.await_count == 1
    items = mem.store_batch.await_args.args[0]
    assert mem.store_batch.await_args.kwargs["collection"] == KNOWLEDGE_COLLECTION == "user_knowledge"
    assert any("Pedraforca-77" in it["text"] for it in items)
    assert all(it["metadata"]["source_document"] == "informe.docx" for it in items)
    assert all(it["metadata"]["session_id"] == "e3-session" for it in items)
    attached_body = session.attach_document.call_args.args[1]
    assert "Pedraforca-77" in attached_body
    # The binary never leaks into the chunks: extracted text, not zip bytes.
    assert not any("PK\x03\x04" in it["text"] or "word/document.xml" in it["text"] for it in items)

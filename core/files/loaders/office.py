"""
────────────────────────────────────
Server Nexe
Author: Jordi Goy
Location: core/files/loaders/office.py
Description: Office and ebook loaders — DOCX, XLSX, PPTX, EPUB.

www.jgoy.net · https://server-nexe.org
────────────────────────────────────

ADR-008 E3 (#1070). All four are zip archives, so every loader here opens the
file through `check_zip` first: the zip magic, a cap on how many members the
archive has and on how much it DECLARES it inflates to. The declared size is a
real bound, not a hint: `zipfile` stops reading a member at its declared size
and fails the CRC if the data goes on, so an archive cannot lie its way past
the cap. Only then does a parser see the file.

Legacy binary formats (`.doc`, `.xls`) and anything needing a system binary
(antiword, tesseract) are deliberately absent. EPUB is read with `zipfile` and
lxml rather than EbookLib, which is AGPL-3.0 and cannot ship in an Apache-2.0
product.
"""

from __future__ import annotations

import logging
import posixpath
import zipfile
from pathlib import Path
from urllib.parse import unquote

from core.files.loaders import (
    MAX_FILE_SIZE,
    ZIP_MAGIC,
    LoaderError,
    TextBudget,
    register_loader,
)
from core.files.loaders.data import safe_xml_parser

logger = logging.getLogger(__name__)

#: Most members an Office/EPUB archive may have. A long novel with its images
#: is a few hundred; ten thousand is not a document, it is a denial of service.
MAX_ZIP_MEMBERS = 10_000

#: Most bytes the archive may declare it inflates to: twenty times the upload
#: cap. XML compresses around 10:1, so a full 10MB spreadsheet fits; a zip bomb
#: (thousands to one) does not.
MAX_ZIP_UNCOMPRESSED = 20 * MAX_FILE_SIZE


def check_zip(file_path: Path) -> None:
    """Refuse a file that is not a zip, or is a zip too big to open safely."""
    with open(file_path, "rb") as fh:
        if fh.read(len(ZIP_MAGIC)) != ZIP_MAGIC:
            raise LoaderError(f"{file_path.name} is not a zip archive")
    try:
        with zipfile.ZipFile(file_path) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as e:
        raise LoaderError(f"{file_path.name}: {e}") from e
    if len(infos) > MAX_ZIP_MEMBERS:
        raise LoaderError(f"{file_path.name}: {len(infos)} archive members (max {MAX_ZIP_MEMBERS})")
    declared = sum(info.file_size for info in infos)
    if declared > MAX_ZIP_UNCOMPRESSED:
        raise LoaderError(
            f"{file_path.name}: inflates to {declared} bytes (max {MAX_ZIP_UNCOMPRESSED})"
        )


_ZIP = [ZIP_MAGIC]


def _row_text(cells) -> str:
    """Cells of one table row, merged cells (repeated by the parsers) once."""
    out: list[str] = []
    for cell in cells:
        text = " ".join(cell.split())
        if text and (not out or out[-1] != text):
            out.append(text)
    return " | ".join(out)


@register_loader([".docx"], magic=_ZIP)
def load_docx_file(file_path: Path) -> str:
    """Paragraphs and tables, in the order they appear in the document."""
    check_zip(file_path)
    from docx import Document
    from docx.table import Table

    doc = Document(str(file_path))
    budget = TextBudget(sep="\n\n")
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            rows = [_row_text(c.text for c in row.cells) for row in block.rows]
            text = "\n".join(r for r in rows if r)
        else:
            text = block.text.strip()
        if text:
            budget.add(text)
    return budget.text()


@register_loader([".xlsx"], magic=_ZIP)
def load_xlsx_file(file_path: Path) -> str:
    """Every sheet, one line per non-empty row, cached values (not formulas)."""
    check_zip(file_path)
    from openpyxl import load_workbook

    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    try:
        budget = TextBudget()
        for ws in wb.worksheets:
            title_added = False
            for row in ws.iter_rows(values_only=True):
                values = ["" if v is None else str(v) for v in row]
                while values and not values[-1].strip():
                    values.pop()
                if not any(v.strip() for v in values):
                    continue
                if not title_added:  # an empty sheet leaves no trace
                    budget.add(("\n" if budget.parts else "") + f"[Sheet: {ws.title}]")
                    title_added = True
                budget.add(" | ".join(values))
        return budget.text()
    finally:
        wb.close()


def _pptx_shape_texts(shape):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for inner in shape.shapes:
            yield from _pptx_shape_texts(inner)
        return
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            yield text
    if shape.has_table:
        rows = [_row_text(c.text for c in row.cells) for row in shape.table.rows]
        text = "\n".join(r for r in rows if r)
        if text:
            yield text


@register_loader([".pptx"], magic=_ZIP)
def load_pptx_file(file_path: Path) -> str:
    """Slide by slide: text boxes, tables, grouped shapes and speaker notes."""
    check_zip(file_path)
    from pptx import Presentation

    prs = Presentation(str(file_path))
    budget = TextBudget(sep="\n\n")
    for number, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            parts.extend(_pptx_shape_texts(shape))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                parts.append(f"Notes: {notes}")
        if parts:
            budget.add(f"[Slide {number}]\n" + "\n".join(parts))
    return budget.text()


# ── EPUB ─────────────────────────────────────────────────────────

_OPF_NS = {"opf": "http://www.idpf.org/2007/opf"}
_CONTAINER_NS = {"c": "urn:oasis:names:tc:opendocument:xmlns:container"}
_XHTML_TYPES = {"application/xhtml+xml", "text/html"}
_BLOCK_TAGS = (
    "p", "div", "li", "tr", "td", "th", "dt", "dd", "blockquote", "pre",
    "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "br",
)


def _epub_spine(zf: zipfile.ZipFile) -> list[str]:
    """Member names of the reading order, from container.xml -> OPF spine.

    Falls back to every (X)HTML member in archive order when the package
    metadata is missing or broken: the text is what matters, not its manifest.
    """
    from lxml import etree

    parser = safe_xml_parser()
    try:
        container = etree.fromstring(zf.read("META-INF/container.xml"), parser)  # nosec B320 - safe parser
        opf_path = container.find(".//c:rootfile", _CONTAINER_NS).get("full-path")
        opf = etree.fromstring(zf.read(opf_path), parser)  # nosec B320 - safe parser
        base = posixpath.dirname(opf_path)
        manifest = {
            item.get("id"): (item.get("href"), item.get("media-type"))
            for item in opf.iterfind(".//opf:manifest/opf:item", _OPF_NS)
        }
        spine = []
        for ref in opf.iterfind(".//opf:spine/opf:itemref", _OPF_NS):
            href, media_type = manifest.get(ref.get("idref"), (None, None))
            if href and media_type in _XHTML_TYPES:
                spine.append(posixpath.normpath(posixpath.join(base, unquote(href))))
        if spine:
            return spine
    except (KeyError, AttributeError, etree.XMLSyntaxError) as e:
        logger.info("EPUB package metadata unreadable (%s); reading every XHTML member", e)
    return [n for n in zf.namelist() if n.lower().endswith((".xhtml", ".html", ".htm"))]


def _xhtml_text(data: bytes) -> str:
    import lxml.html

    # EPUB content documents are UTF-8 or UTF-16 (EPUB 3 §3.1); without saying
    # so, lxml's HTML parser assumes latin-1 and "Capítol" becomes "CapÃ­tol".
    encoding = "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8"
    parser = lxml.html.HTMLParser(
        encoding=encoding, no_network=True, remove_comments=True, remove_pis=True,
    )
    root = lxml.html.document_fromstring(data, parser=parser)
    for bad in list(root.iter("script", "style", "head")):
        bad.drop_tree()
    for el in root.iter(*_BLOCK_TAGS):
        el.tail = "\n" + (el.tail or "")
    lines = (" ".join(line.split()) for line in root.text_content().splitlines())
    return "\n".join(line for line in lines if line)


@register_loader([".epub"], magic=_ZIP)
def load_epub_file(file_path: Path) -> str:
    """The chapters' text in reading order."""
    check_zip(file_path)
    budget = TextBudget(sep="\n\n")
    with zipfile.ZipFile(file_path) as zf:
        names = set(zf.namelist())
        for name in _epub_spine(zf):
            if name not in names:
                continue
            text = _xhtml_text(zf.read(name))
            if text:
                budget.add(text)
    return budget.text()
